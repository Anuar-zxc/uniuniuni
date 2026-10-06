"""Secure CV file validation and text extraction (PDF, DOCX, TXT)."""

import io
import re

ALLOWED = {
    ".pdf": ("application/pdf", b"%PDF"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04"),
    ".txt": ("text/plain", None),
}


class ResumeParseError(ValueError):
    pass


def validate_upload(filename: str, data: bytes, max_mb: int) -> str:
    name = (filename or "").lower()
    ext = next((e for e in ALLOWED if name.endswith(e)), None)
    if not ext:
        raise ResumeParseError("Only PDF, DOCX or TXT files are supported")
    if len(data) == 0:
        raise ResumeParseError("File is empty")
    if len(data) > max_mb * 1024 * 1024:
        raise ResumeParseError(f"File is larger than {max_mb} MB")
    magic = ALLOWED[ext][1]
    if magic and not data.startswith(magic):
        raise ResumeParseError("File content does not match its extension")
    return ext


def extract_text(ext: str, data: bytes) -> str:
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeParseError("Encrypted PDFs are not supported")
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:15])
        elif ext == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="ignore")
    except ResumeParseError:
        raise
    except Exception as e:
        raise ResumeParseError("Could not read the file") from e
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise ResumeParseError("Could not extract enough text (is it a scanned image?). Try DOCX or paste text.")
    return text[:30000]
